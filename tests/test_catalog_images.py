"""The catalog template, product photos, and the merchant-facing error summary.

Column mapping is deterministic now - the template ships the headers the alias table knows -
so the mapping tests here are about that contract holding, not about a model. The model is
still used to screen products and to match leftover photo filenames; its contribution is
tested where it matters, at the validators that decide whether its answer may touch anything.

The suite runs in DEMO_MODE, so model calls are skipped and the deterministic path executes.
"""

from __future__ import annotations

import asyncio
import io
import zipfile
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.db import connect, init_databases
from app.main import app
from merchant import catalog_images
from merchant.catalog_diagnostics import (
    DiagnosticsValidationError,
    build_groups,
    classify_issue,
    deterministic_headline,
    validate_diagnostics_payload,
)
from merchant.catalog_images import (
    ImageArchiveError,
    ImageMatchValidationError,
    ImageTarget,
    extract_image_archive,
    match_images,
    normalize_name,
    validate_match_payload,
)
from merchant.catalog_mapping import (
    FIELD_ALIASES,
    MODEL_DESCRIPTIVE_FIELDS,
    detect_mappings,
    resolve_mappings,
)
from merchant.catalog_pipeline import SourceRow, _descriptive_fields
from merchant.catalog_template import COLUMNS
from seed.reset import MERCHANT_KEY_FILE, seed

PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00"
    b"\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05\x00\x01\r\n-\xb4"
    b"\x00\x00\x00\x00IEND\xaeB`\x82"
)


@pytest.fixture()
def client() -> TestClient:
    init_databases(reset=True)
    seed()
    with TestClient(app) as test_client:
        test_client.headers["X-Merchant-Key"] = MERCHANT_KEY_FILE.read_text(
            encoding="utf-8"
        ).strip()
        yield test_client


def make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def workbook_bytes(rows: list[list[object]], headers: list[str]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Products"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


CATALOG_HEADERS = ["sku", "title", "price", "ingredients", "stock", "product_type", "description"]
CATALOG_ROWS = [
    ["MYSA-01", "Gentle Cloud Cleanser", 29.90, "glycerin, panthenol", 10, "face wash",
     "A gentle hydrating face wash for dry skin"],
    ["MYSA-02", "Barrier Milk Moisturiser", 42.00, "ceramide, squalane", 6, "moisturiser",
     "A soothing cream for barrier support"],
    ["MYSA-03", "Calm Gel Serum", 35.00, "niacinamide", 4, "serum",
     "A light serum for oily skin and redness"],
]


def publish_catalog(client: TestClient, rows=None, headers=None) -> dict:
    """Upload and publish, so there are live products for photos to attach to."""
    upload = client.post(
        "/merchant/m_mysa/catalog/uploads",
        files={"file": ("catalog.xlsx", workbook_bytes(rows or CATALOG_ROWS, headers or CATALOG_HEADERS),
                        "application/vnd.ms-excel")},
    ).json()
    plan = upload["approval"]["modes"]["replace"]
    approved = client.post(
        f"/merchant/m_mysa/catalog/uploads/{upload['upload_id']}/approve",
        json={
            "approval_token": plan["approval_token"],
            "reviewed_row_count": upload["approval"]["reviewed_row_count_required"],
            "mode": "replace",
        },
    )
    assert approved.status_code == 200, approved.text
    return upload


# --- the catalog template ----------------------------------------------------------------


def test_template_downloads_as_a_workbook(client: TestClient) -> None:
    response = client.get("/catalog/template")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "skincare-catalog-template.xlsx" in response.headers["content-disposition"]

    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Products", "How to fill this in"]
    headers = [cell.value for cell in next(book["Products"].iter_rows(max_row=1))]
    assert headers == [column for column, _, _, _ in COLUMNS]


def test_every_template_header_is_a_header_the_importer_knows() -> None:
    """The template's whole purpose: its headers map with no guessing at all."""
    mapped = {alias for aliases in FIELD_ALIASES.values() for alias in aliases}
    for column, _, _, _ in COLUMNS:
        assert column in mapped or column in MODEL_DESCRIPTIVE_FIELDS, (
            f"template column {column!r} is neither a mapped field nor recognised evidence, "
            "so filling it in would have no effect"
        )


def test_a_file_filled_in_from_the_template_maps_perfectly(client: TestClient) -> None:
    template = client.get("/catalog/template").content
    book = load_workbook(io.BytesIO(template))
    sheet = book["Products"]
    # The merchant deletes the example rows and types their own, as the guide instructs.
    sheet.delete_rows(2, sheet.max_row)
    sheet.append(["ACME-1", "Quiet Balm Cleanser", 24.50, 8, "aqua, glycerin",
                  "A calming balm cleanser.", "cleansing balm", "dry|sensitive", "dryness",
                  "evening", "balm", "yes", "", 100, "", "", "", ""])
    buffer = io.BytesIO()
    book.save(buffer)

    upload = client.post(
        "/merchant/m_mysa/catalog/uploads",
        files={"file": ("filled.xlsx", buffer.getvalue(), "application/vnd.ms-excel")},
    ).json()

    assert upload["summary"]["ready"] == 1, upload["errors"]
    report = upload["mapping_report"]
    assert report["source"] == "deterministic_alias_table"
    assert report["unresolved"] == []
    assert report["ignored_columns"] == []
    assert {d["method"] for d in report["decisions"]} == {"exact_alias"}
    canonical = upload["products"][0]["canonical"]
    assert canonical["sku"] == "ACME-1"
    assert canonical["attributes"]["product_type"] == "cleansing_balm"
    assert set(canonical["attributes"]["skin_types"]) == {"dry", "sensitive"}


def test_optional_template_columns_may_be_left_blank(client: TestClient) -> None:
    """Only the five required columns must be filled; the rest are evidence when present."""
    required_only = ["sku", "title", "price", "stock", "ingredients"]
    upload = client.post(
        "/merchant/m_mysa/catalog/uploads",
        files={
            "file": (
                "sparse.xlsx",
                workbook_bytes([["BARE-1", "Plain Cleanser", 19.00, 3, "aqua, glycerin"]],
                               required_only),
                "application/vnd.ms-excel",
            )
        },
    ).json()
    assert upload["summary"]["rejected"] == 0, upload["errors"]
    assert upload["products"][0]["canonical"]["sku"] == "BARE-1"


# --- deterministic column mapping ---------------------------------------------------------


def rows_with(headers: list[str]) -> list[SourceRow]:
    return [
        SourceRow(
            source_record_id="row_000002",
            row_number=2,
            sheet_name="Products",
            values={header: f"value for {header}" for header in headers},
        )
    ]


def test_alias_ties_are_recorded_then_stop_the_upload(client: TestClient) -> None:
    candidates = detect_mappings(rows_with(["sku", "price", "unit_price", "stock"]))
    assert candidates.conflicts["price"] == ["price", "unit_price"]

    response = client.post(
        "/merchant/m_mysa/catalog/uploads",
        files={
            "file": (
                "ambiguous.xlsx",
                workbook_bytes([["A-1", "Serum", 10, 10, "niacinamide", 5]],
                               ["sku", "title", "price", "unit_price", "ingredients", "stock"]),
                "application/vnd.ms-excel",
            )
        },
    )
    assert response.status_code == 400
    error = response.json()["detail"]["error"]
    assert error["details"]["unresolved_fields"][0]["candidate_columns"] == ["price", "unit_price"]
    assert "catalog template" in error["message"]


def test_unrecognised_columns_are_reported_not_guessed(client: TestClient) -> None:
    headers = [*CATALOG_HEADERS, "Landed cost", "Supplier ref"]
    rows = [[*CATALOG_ROWS[0], 8.10, "SUP-1"]]
    upload = client.post(
        "/merchant/m_mysa/catalog/uploads",
        files={"file": ("extra.xlsx", workbook_bytes(rows, headers), "application/vnd.ms-excel")},
    ).json()

    assert upload["mapping_report"]["ignored_columns"] == ["Landed cost", "Supplier ref"]
    assert upload["summary"]["ready"] == 1
    assert any("not recognised" in note for note in upload["diagnostics"]["notes"])


def test_mapping_never_calls_a_model() -> None:
    """resolve_mappings is synchronous by construction - there is no call to make."""
    resolution = resolve_mappings(detect_mappings(rows_with(["sku", "title", "mystery"])))
    assert resolution.source == "deterministic_alias_table"
    assert resolution.mappings == {"sku": "sku", "title": "title"}
    assert resolution.ignored_columns == ["mystery"]


def test_a_mapped_column_reaches_the_screener_but_a_locked_fact_never_does() -> None:
    row = SourceRow(
        source_record_id="row_000002",
        row_number=2,
        sheet_name=None,
        values={"sku": "A-1", "title": "Calm Serum", "price": "30.00",
                "full_inci": "niacinamide", "Landed cost": "8.00"},
    )
    fields = _descriptive_fields(
        row, {"sku": "sku", "title": "title", "price": "price", "ingredients": "full_inci"}
    )
    assert fields == {"title": "Calm Serum", "ingredients": "niacinamide"}


# --- photos on the live catalog -----------------------------------------------------------


def test_photos_attach_to_live_products_straight_away(client: TestClient) -> None:
    publish_catalog(client)
    response = client.post(
        "/merchant/m_mysa/catalog/images",
        files={
            "file": (
                "photos.zip",
                make_zip({
                    "MYSA-01.png": PNG,                    # matches the SKU
                    "barrier milk moisturiser.png": PNG,   # matches the product name
                    "Calm-Gel-Serum.png": PNG,             # matches once separators are folded
                    "not-a-product.png": PNG,              # matches nothing
                }),
                "application/zip",
            )
        },
    )
    assert response.status_code == 200, response.text
    report = response.json()

    assert report["image_count"] == 4
    assert report["matched_count"] == 3
    assert report["product_count"] == 3
    assert report["unmatched_images"] == ["not-a-product.png"]
    assert report["match_source"] == "deterministic"
    # An unmatched photo is reported but never stored, so it cannot leave a dead link.
    unmatched = next(i for i in report["images"] if not i["matched"])
    assert unmatched["url"] is None
    with connect() as connection:
        assert connection.execute("SELECT COUNT(*) c FROM catalog_images").fetchone()["c"] == 3

    methods = {i["entry_name"]: i.get("method") for i in report["images"] if i["matched"]}
    assert methods == {
        "MYSA-01.png": "exact_sku",
        "barrier milk moisturiser.png": "exact_name",
        "Calm-Gel-Serum.png": "exact_name",
    }

    # The live product is updated immediately - no second approval step.
    product = client.get("/catalog/product/MYSA-01?merchant_id=m_mysa").json()
    assert product["image_url"].startswith("/api/catalog/images/img_")


def test_photos_are_refused_when_there_is_no_catalog_yet(client: TestClient) -> None:
    other = client.post("/merchant/onboard", json={"name": "Fresh Skin", "size": "sme"}).json()
    fresh = TestClient(app)
    fresh.headers["X-Merchant-Key"] = other["api_key"]
    response = fresh.post(
        f"/merchant/{other['merchant_id']}/catalog/images",
        files={"file": ("photos.zip", make_zip({"A-1.png": PNG}), "application/zip")},
    )
    assert response.status_code == 409
    assert response.json()["detail"]["error"]["code"] == "NO_PRODUCTS"


def test_photos_only_ever_change_the_picture(client: TestClient) -> None:
    publish_catalog(client)
    before = client.get("/catalog/product/MYSA-01?merchant_id=m_mysa").json()
    client.post(
        "/merchant/m_mysa/catalog/images",
        files={"file": ("photos.zip", make_zip({"MYSA-01.png": PNG}), "application/zip")},
    )
    after = client.get("/catalog/product/MYSA-01?merchant_id=m_mysa").json()

    assert after["image_url"] != before["image_url"]
    for field in ("sku", "title", "description", "price_cents", "currency", "stock",
                  "rating_avg", "rating_count", "attributes"):
        assert after[field] == before[field], f"{field} changed"


def test_a_corrected_archive_replaces_the_photo_and_drops_the_old_bytes(client: TestClient) -> None:
    publish_catalog(client)
    client.post(
        "/merchant/m_mysa/catalog/images",
        files={"file": ("a.zip", make_zip({"MYSA-01.png": PNG}), "application/zip")},
    )
    first = client.get("/catalog/product/MYSA-01?merchant_id=m_mysa").json()["image_url"]

    client.post(
        "/merchant/m_mysa/catalog/images",
        files={"file": ("b.zip", make_zip({"MYSA-01.png": PNG}), "application/zip")},
    )
    second = client.get("/catalog/product/MYSA-01?merchant_id=m_mysa").json()["image_url"]
    assert second != first

    with connect() as connection:
        stored = connection.execute("SELECT image_id FROM catalog_images").fetchall()
    assert [row["image_id"] for row in stored] == [second.rsplit("/", 1)[-1]]


def test_photos_are_merchant_only_until_the_store_is_published(client: TestClient) -> None:
    publish_catalog(client)
    client.put("/merchant/m_mysa/config", json={"status": "draft"})
    report = client.post(
        "/merchant/m_mysa/catalog/images",
        files={"file": ("photos.zip", make_zip({"MYSA-01.png": PNG}), "application/zip")},
    ).json()
    path = report["images"][0]["url"].removeprefix("/api")

    anonymous = TestClient(app)
    assert anonymous.get(path).status_code == 401
    assert client.get(path).status_code == 200

    client.put("/merchant/m_mysa/config", json={"status": "published"})
    live = anonymous.get(path)
    assert live.status_code == 200
    assert live.headers["content-type"] == "image/png"
    assert live.headers["x-content-type-options"] == "nosniff"
    assert live.content == PNG


def test_one_merchants_photos_never_reach_another(client: TestClient) -> None:
    publish_catalog(client)
    other = client.post("/merchant/onboard", json={"name": "Rival Skin", "size": "sme"}).json()
    intruder = TestClient(app)
    intruder.headers["X-Merchant-Key"] = other["api_key"]
    client.post(
        "/merchant/m_mysa/catalog/images",
        files={"file": ("photos.zip", make_zip({"MYSA-01.png": PNG}), "application/zip")},
    )
    mine = client.get("/catalog/product/MYSA-01?merchant_id=m_mysa").json()["image_url"]
    assert mine is not None

    # A rival naming the same SKU has no products of their own for it to hit, and cannot
    # reach into Mysa's catalog to repoint a picture.
    response = intruder.post(
        f"/merchant/{other['merchant_id']}/catalog/images",
        files={"file": ("photos.zip", make_zip({"MYSA-01.png": PNG}), "application/zip")},
    )
    assert response.status_code == 409
    assert client.get("/catalog/product/MYSA-01?merchant_id=m_mysa").json()["image_url"] == mine


# --- archive hardening -------------------------------------------------------------------


@pytest.mark.parametrize(
    ("entries", "kept", "reason_fragment"),
    [
        ({"ok.png": PNG, "../escape.png": PNG}, ["ok.png"], "unsafe"),
        ({"ok.png": PNG, "notes.txt": b"hello"}, ["ok.png"], "extension"),
        ({"ok.png": PNG, "fake.png": b"MZ this is an executable"}, ["ok.png"], "not a known image"),
        ({"ok.png": PNG, "__MACOSX/ok.png": PNG}, ["ok.png"], "unsafe"),
        ({"ok.png": PNG, "OK.PNG": PNG}, ["ok.png"], "duplicate product name"),
    ],
)
def test_archive_refuses_unsafe_entries_and_says_why(
    entries: dict[str, bytes], kept: list[str], reason_fragment: str
) -> None:
    images, skipped = extract_image_archive(make_zip(entries), "images.zip")
    assert [image.entry_name for image in images] == kept
    assert any(reason_fragment in item["reason"] for item in skipped), skipped


def test_archive_rejects_non_zip_and_empty_archives() -> None:
    with pytest.raises(ImageArchiveError, match="valid ZIP"):
        extract_image_archive(b"not a zip at all", "images.zip")
    with pytest.raises(ImageArchiveError, match=r"\.zip archive"):
        extract_image_archive(make_zip({"ok.png": PNG}), "images.rar")
    with pytest.raises(ImageArchiveError, match="no usable image"):
        extract_image_archive(make_zip({"notes.txt": b"hello"}), "images.zip")


def test_oversized_archive_is_refused_by_the_route(client: TestClient) -> None:
    publish_catalog(client)
    response = client.post(
        "/merchant/m_mysa/catalog/images",
        files={"file": ("photos.zip", b"x" * (25 * 1024 * 1024 + 1), "application/zip")},
    )
    assert response.status_code == 413


# --- image_url from the workbook ---------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("https://cdn.example.com/a.png", "https://cdn.example.com/a.png"),
        ("/products/a.png", "/products/a.png"),
        ("", None),
    ],
)
def test_workbook_image_urls_that_are_safe_to_render(
    client: TestClient, value: str, expected: str | None
) -> None:
    headers = [*CATALOG_HEADERS, "image_url"]
    upload = client.post(
        "/merchant/m_mysa/catalog/uploads",
        files={"file": ("catalog.xlsx", workbook_bytes([[*CATALOG_ROWS[0], value]], headers),
                        "application/vnd.ms-excel")},
    ).json()
    assert upload["summary"]["ready"] == 1
    assert upload["products"][0]["canonical"]["image_url"] == expected


@pytest.mark.parametrize(
    "value",
    ["javascript:alert(1)", "data:image/png;base64,AAAA", "//evil.example/a.png", "ftp://h/a.png"],
)
def test_workbook_image_urls_that_cannot_be_rendered_are_held(
    client: TestClient, value: str
) -> None:
    headers = [*CATALOG_HEADERS, "image_url"]
    upload = client.post(
        "/merchant/m_mysa/catalog/uploads",
        files={"file": ("catalog.xlsx", workbook_bytes([[*CATALOG_ROWS[0], value]], headers),
                        "application/vnd.ms-excel")},
    ).json()
    assert upload["summary"]["rejected"] == 1
    assert "image_url must be an https:// link" in upload["errors"][0]["reason"]


# --- matching internals ------------------------------------------------------------------


def test_fuzzy_matching_needs_a_clear_winner() -> None:
    targets = [
        ImageTarget("A-1", 1, "A-1", "Calm Gel Serum"),
        ImageTarget("A-2", 2, "A-2", "Calm Gel Cream"),
    ]
    images, _ = extract_image_archive(make_zip({"calm gel serum.png": PNG}), "i.zip")
    matched, _ = asyncio.run(match_images(images, targets))
    assert matched["calm gel serum.png"]["source_record_id"] == "A-1"

    ambiguous, _ = extract_image_archive(make_zip({"calm gel.png": PNG}), "i.zip")
    assert asyncio.run(match_images(ambiguous, targets))[0] == {}


def test_normalize_name_folds_separators_and_case() -> None:
    assert normalize_name("Calm-Gel_Serum  30ml.PNG") == "calm gel serum 30ml png"
    assert normalize_name("MYSA--01") == "mysa 01"


def test_model_image_matches_must_name_real_ids_and_stay_one_to_one() -> None:
    ids, records = {"a.png", "b.png"}, {"A-1", "A-2"}
    accepted = validate_match_payload(
        {
            "schema_version": "catalog-image-match.v1",
            "matches": [
                {"image_id": "a.png", "source_record_id": "A-1", "confidence": 0.9, "reason": "same name"},
                {"image_id": "b.png", "source_record_id": "A-2", "confidence": 0.4, "reason": "unsure"},
            ],
        },
        image_ids=ids,
        record_ids=records,
    )
    assert [item["image_id"] for item in accepted] == ["a.png"]  # 0.4 is below the floor

    for payload, message in [
        ([{"image_id": "ghost.png", "source_record_id": "A-1", "confidence": 1, "reason": "x"}],
         "unknown image"),
        ([{"image_id": "a.png", "source_record_id": "NOPE", "confidence": 1, "reason": "x"}],
         "unknown product row"),
        ([{"image_id": "a.png", "source_record_id": "A-1", "confidence": 1, "reason": "x"},
          {"image_id": "a.png", "source_record_id": "A-2", "confidence": 1, "reason": "x"}],
         "matched twice"),
    ]:
        with pytest.raises(ImageMatchValidationError, match=message):
            validate_match_payload(
                {"schema_version": "catalog-image-match.v1", "matches": payload},
                image_ids=ids,
                record_ids=records,
            )
    with pytest.raises(ImageMatchValidationError, match="schema_version"):
        validate_match_payload({"schema_version": "wrong.v9", "matches": []},
                               image_ids=ids, record_ids=records)


def live_model(monkeypatch, module) -> None:
    """Pretend this process has a usable model, without one being reachable."""
    monkeypatch.setattr(
        module,
        "settings",
        SimpleNamespace(demo_mode=False, openai_api_key="test", openai_model="test-model"),
    )


def test_the_model_only_sees_photos_the_deterministic_pass_could_not_place(monkeypatch) -> None:
    live_model(monkeypatch, catalog_images)
    targets = [ImageTarget("A-1", 1, "A-1", "Calm Gel Serum"),
               ImageTarget("A-2", 2, "A-2", "Barrier Milk")]
    images, _ = extract_image_archive(
        make_zip({"A-1.png": PNG, "barrier-milk-50ml-new.png": PNG}), "i.zip"
    )

    async def fake_ask(unmatched_images, unmatched_targets):
        assert [i.entry_name for i in unmatched_images] == ["barrier-milk-50ml-new.png"]
        assert [t.sku for t in unmatched_targets] == ["A-2"]
        return [{"image_id": "barrier-milk-50ml-new.png", "source_record_id": "A-2",
                 "confidence": 0.88, "reason": "same product with a size suffix"}]

    monkeypatch.setattr(catalog_images, "_ask_model", fake_ask)
    matched, source = asyncio.run(catalog_images.match_images(images, targets))
    assert matched["A-1.png"]["method"] == "exact_sku"
    assert matched["barrier-milk-50ml-new.png"]["method"] == "model"
    assert source == "hybrid_model_with_deterministic_guard"


def test_a_photo_model_failure_keeps_the_deterministic_matches(monkeypatch) -> None:
    live_model(monkeypatch, catalog_images)

    async def explode(*_args):
        raise RuntimeError("synthetic match failure")

    monkeypatch.setattr(catalog_images, "_ask_model", explode)
    targets = [ImageTarget("A-1", 1, "A-1", "Calm Gel Serum"),
               ImageTarget("A-2", 2, "A-2", "Barrier Milk")]
    images, _ = extract_image_archive(make_zip({"A-1.png": PNG, "mystery.png": PNG}), "i.zip")
    matched, source = asyncio.run(catalog_images.match_images(images, targets))
    assert set(matched) == {"A-1.png"}
    assert source == "deterministic_failover"


# --- merchant-facing error summary -------------------------------------------------------


def test_upload_diagnostics_explain_every_held_row(client: TestClient) -> None:
    rows = [
        ["OK-1", "Gentle Cloud Cleanser", 29.90, "glycerin", 10, "face wash", "A gentle face wash for dry skin"],
        ["BAD-1", "Broken Price Toner", "N/A", "glycerin", 3, "toner", "A toner"],
        ["BAD-2", "No Ingredients Serum", 20.00, "", 3, "serum", "A serum"],
        ["OK-1", "Duplicate Code Cream", 25.00, "ceramide", 3, "moisturiser", "A moisturiser"],
    ]
    upload = client.post(
        "/merchant/m_mysa/catalog/uploads",
        files={"file": ("catalog.xlsx", workbook_bytes(rows, CATALOG_HEADERS),
                        "application/vnd.ms-excel")},
    ).json()

    diagnostics = upload["diagnostics"]
    assert diagnostics["source"] == "deterministic"
    assert diagnostics["headline"] == "1 of 4 rows are ready to publish; 3 could not be read."
    codes = {group["code"]: group for group in diagnostics["groups"]}
    assert set(codes) == {"PRICE_INVALID", "MISSING_INGREDIENTS", "DUPLICATE_SKU"}
    assert codes["DUPLICATE_SKU"]["row_count"] == 1
    assert codes["DUPLICATE_SKU"]["example_rows"] == [5]
    assert all(group["blocking"] for group in diagnostics["groups"])


def test_a_clean_upload_says_so(client: TestClient) -> None:
    upload = client.post(
        "/merchant/m_mysa/catalog/uploads",
        files={"file": ("catalog.xlsx", workbook_bytes(CATALOG_ROWS, CATALOG_HEADERS),
                        "application/vnd.ms-excel")},
    ).json()
    assert upload["diagnostics"]["groups"] == []
    assert upload["diagnostics"]["headline"] == "All 3 rows are clean and ready to publish."


@pytest.mark.parametrize(
    ("issue", "code"),
    [
        ("price must be numeric with at most two decimal places", "PRICE_INVALID"),
        ("stock is required", "STOCK_INVALID"),
        ("duplicate sku A-1; each source row must be unique", "DUPLICATE_SKU"),
        ("formula content requires review: Notes", "FORMULA_CELL"),
        ("image_url must be an https:// link or a path on this store", "IMAGE_URL_INVALID"),
        ("exactly one primary product type is required", "AMBIGUOUS_PRODUCT_TYPE"),
        ("something nobody planned for", "OTHER"),
    ],
)
def test_issues_are_grouped_under_a_stable_code(issue: str, code: str) -> None:
    assert classify_issue(issue) == code


def test_one_row_with_two_problems_is_counted_once_per_group() -> None:
    groups = build_groups(
        [
            {"row_number": 2, "status": "rejected", "issues": ["sku is required", "price is required"]},
            {"row_number": 3, "status": "rejected", "issues": ["price is required"]},
            {"row_number": 4, "status": "ready", "issues": []},
        ]
    )
    counts = {group["code"]: group["row_count"] for group in groups}
    assert counts == {"PRICE_INVALID": 2, "MISSING_SKU": 1}
    assert deterministic_headline(
        {"input_rows": 3, "ready": 1, "review_required": 0, "rejected": 2}, groups
    ) == "1 of 3 rows are ready to publish; 2 could not be read."


def test_the_model_may_reword_a_group_but_never_invent_one() -> None:
    rewritten = validate_diagnostics_payload(
        {
            "schema_version": "catalog-diagnostics.v1",
            "headline": "Most of your file is ready.",
            "groups": [
                {"code": "PRICE_INVALID", "title": "Prices we could not read",
                 "why": "The price cell was not a number.", "fix": "Write prices as 29.90."}
            ],
        },
        codes=["PRICE_INVALID"],
    )
    assert rewritten["prose"]["PRICE_INVALID"]["title"] == "Prices we could not read"

    for payload, message in [
        ({"schema_version": "catalog-diagnostics.v1", "headline": "h", "groups": [
            {"code": "MADE_UP", "title": "t", "why": "w", "fix": "f"}]}, "exactly the codes"),
        ({"schema_version": "catalog-diagnostics.v1", "headline": "h", "groups": []},
         "exactly the codes"),
        ({"schema_version": "old.v0", "headline": "h", "groups": []}, "schema_version"),
    ]:
        with pytest.raises(DiagnosticsValidationError, match=message):
            validate_diagnostics_payload(payload, codes=["PRICE_INVALID"])
